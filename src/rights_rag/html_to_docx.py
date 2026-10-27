"""Convert Kol Zchut pages saved from Chrome as HTML into Word documents.

The goal is a realistic Word file, not a clean one: nothing is filtered here
(that is the loader's job, as with PDF). Headings become Word heading styles,
lists become list paragraphs, tables become real Word tables, and everything
is right-to-left.

    uv run python -m rights_rag.html_to_docx docs/html docs/docx
"""

import re
import sys
from pathlib import Path

from bs4 import BeautifulSoup, Comment, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

HEADING_STYLES = {
    "h2": "Heading 1",
    "h3": "Heading 2",
    "h4": "Heading 3",
    "h5": "Heading 3",
}
INLINE = {
    "a", "abbr", "b", "bdi", "bdo", "big", "cite", "code", "em", "font", "i",
    "label", "mark", "q", "s", "small", "span", "strong", "sub", "sup", "time", "u",
}  # fmt: skip
BLOCK = {"div", "p", "ul", "ol", "table", "h2", "h3", "h4", "h5", "dl", "blockquote"}
SKIP = {"style", "script", "link", "meta", "img", "svg", "noscript", "button", "input"}


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _rtl(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
    for run in paragraph.runs:
        run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))


def _heading_text(el: Tag) -> str:
    headline = el.select_one(".mw-headline")
    if headline:
        return _norm(headline.get_text())
    for edit in el.select(".mw-editsection"):
        edit.decompose()
    return _norm(el.get_text())


class _Writer:
    def __init__(self, doc):
        self.doc = doc
        self.buffer: list[str] = []

    def para(self, text: str, style: str = "Normal", container=None) -> None:
        text = _norm(text)
        if text:
            _rtl((container or self.doc).add_paragraph(text, style=style))

    def flush(self) -> None:
        self.para("".join(self.buffer))
        self.buffer = []

    def walk(self, el: Tag) -> None:
        for child in el.children:
            if isinstance(child, Comment):
                continue
            if isinstance(child, NavigableString):
                self.buffer.append(str(child))
                continue
            name = child.name
            if name in SKIP:
                continue
            if name == "br":
                self.flush()
            elif name in INLINE and not child.find(BLOCK):
                self.buffer.append(child.get_text())
            elif name in HEADING_STYLES:
                self.flush()
                self.para(_heading_text(child), HEADING_STYLES[name])
            elif name == "p":
                self.flush()
                self.walk(child)
                self.flush()
            elif name in ("ul", "ol"):
                self.flush()
                self.list(child, level=1)
            elif name == "table":
                self.flush()
                self.table(child)
            else:
                self.flush()
                self.walk(child)
                self.flush()

    def list(self, el: Tag, level: int, container=None) -> None:
        base = "List Number" if el.name == "ol" else "List Bullet"
        style = base if level == 1 else f"{base} {min(level, 3)}"
        for li in el.find_all("li", recursive=False):
            nested = li.find_all(["ul", "ol"], recursive=False)
            for sub in nested:
                sub.extract()
            self.para(li.get_text(), style, container)
            for sub in nested:
                self.list(sub, level + 1, container)

    def table(self, el: Tag) -> None:
        rows = [tr.find_all(["th", "td"], recursive=False) for tr in el.find_all("tr")]
        rows = [r for r in rows if r]
        if not rows:
            return
        n_cols = max(sum(int(c.get("colspan", 1)) for c in r) for r in rows)
        table = self.doc.add_table(rows=len(rows), cols=n_cols)
        table.style = "Table Grid"
        table._tbl.tblPr.append(OxmlElement("w:bidiVisual"))
        for i, cells in enumerate(rows):
            if all(c.name == "th" for c in cells):
                table.rows[i]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            col = 0
            for cell in cells:
                span = int(cell.get("colspan", 1))
                target = table.cell(i, col)
                if span > 1:
                    target = target.merge(table.cell(i, col + span - 1))
                self._fill_cell(target, cell)
                col += span

    def _fill_cell(self, target, cell: Tag) -> None:
        target.paragraphs[0]._p.getparent().remove(target.paragraphs[0]._p)
        lists = cell.find_all(["ul", "ol"])
        for lst in lists:
            lst.extract()
        self.para(cell.get_text(" "), container=target)
        for lst in lists:
            self.list(lst, level=1, container=target)
        if not target.paragraphs:
            target.add_paragraph("")


def convert(html_path: Path, docx_path: Path) -> None:
    soup = BeautifulSoup(html_path.read_text(encoding="utf-8"), "html.parser")
    doc = Document()
    doc.core_properties.title = (
        _norm(soup.title.get_text()) if soup.title else html_path.stem
    )
    doc.core_properties.language = "he-IL"
    heading = soup.select_one("#firstHeading")
    if heading:
        _rtl(doc.add_paragraph(_norm(heading.get_text()), style="Title"))
    writer = _Writer(doc)
    writer.walk(soup.select_one(".mw-parser-output"))
    writer.flush()
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)


def docx_name(html_path: Path) -> str:
    return html_path.with_suffix(".docx").name


def main(src: str = "docs/html", dst: str = "docs/docx") -> None:
    paths = sorted(Path(src).glob("*.html"))
    for path in paths:
        convert(path, Path(dst) / docx_name(path))
    print(f"converted {len(paths)} files into {dst}")


if __name__ == "__main__":
    main(*sys.argv[1:])
